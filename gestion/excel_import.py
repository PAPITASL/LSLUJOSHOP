from decimal import Decimal, InvalidOperation
from io import BytesIO
import re
import unicodedata

from django.core.exceptions import ValidationError
from django.db import transaction
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from .models import Clientes, Productos, Vendedores


IMPORT_CONFIG = {
    "productos": {
        "model": Productos,
        "columns": (
            ("id", "ID (opcional)"), ("nombre", "Nombre *"), ("categoria", "Categoría *"), ("marca", "Marca"),
            ("modelo", "Modelo"), ("anio_inicio", "Año inicio"), ("anio_fin", "Año fin"),
            ("descripcion", "Descripción"), ("cantidad_disponible", "Cantidad disponible *"),
            ("costo_dolares", "Costo USD"), ("costo_pesos", "Costo COP"),
            ("precio_venta_sugerido", "Precio de venta COP *"), ("estado", "Estado *"),
        ),
        "required": {"nombre", "categoria", "cantidad_disponible", "precio_venta_sugerido", "estado"},
        "defaults": {"costo_dolares": Decimal("0"), "costo_pesos": Decimal("0")},
        "decimal": {"costo_dolares", "costo_pesos", "precio_venta_sugerido"},
        "integer": {"anio_inicio", "anio_fin", "cantidad_disponible"},
        "choices": {
            "categoria": ("FAROLAS", "STOPS", "PERSIANAS / PARRILLAS", "EXPLORADORAS", "LUCES Y DIRECCIONALES", "CARROCERÍA", "INTERIOR", "ELÉCTRICO", "SUSPENSIÓN", "ESCAPE", "EMBLEMAS Y ACCESORIOS", "OTROS"),
            "estado": ("Disponible", "Reservado", "Vendido", "Por pedir", "Agotado"),
        },
        "example": (None, "Farola delantera", "Farolas", "Ford", "F-150", 2015, 2020, "Farola derecha", 4, 50, 200000, 350000, "Disponible"),
    },
    "clientes": {
        "model": Clientes,
        "columns": (("id", "ID (opcional)"), ("nombre", "Nombre *"), ("telefono", "Teléfono"), ("ciudad", "Ciudad"), ("direccion", "Dirección"), ("observaciones", "Observaciones")),
        "required": {"nombre"}, "decimal": set(), "integer": set(), "choices": {}, "defaults": {"telefono": "Sin registrar"},
        "example": (None, "David Calero", "321591172", "Guacarí", "Calle 3 N.º 2-154", "Cliente frecuente"),
    },
    "vendedores": {
        "model": Vendedores,
        "columns": (("id", "ID (opcional)"), ("nombre", "Nombre *"), ("telefono", "Teléfono"), ("porcentaje_comision", "Comisión % *"), ("activo", "Activo *"), ("observaciones", "Observaciones")),
        "required": {"nombre", "porcentaje_comision", "activo"}, "decimal": {"porcentaje_comision"}, "integer": set(), "defaults": {},
        "choices": {"activo": ("Sí", "No")},
        "example": (None, "Vendedor ejemplo", "3001234567", 10, "Sí", ""),
    },
}


def build_template(module):
    config = IMPORT_CONFIG[module]
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = config["model"]._meta.verbose_name_plural.title()
    headers = [label for _field, label in config["columns"]]
    sheet.append(headers)
    sheet.append(config["example"])
    red_fill = PatternFill("solid", fgColor="E6313A")
    for cell in sheet[1]:
        cell.fill = red_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center")
    for index, (_field, label) in enumerate(config["columns"], 1):
        sheet.column_dimensions[chr(64 + index)].width = max(15, min(35, len(label) + 5))
    if module == "productos":
        column_by_field = {field: index for index, (field, _label) in enumerate(config["columns"], 1)}
        for row in range(2, 5001):
            sheet.cell(row, column_by_field["costo_dolares"]).number_format = '$#,##0.00'
            for field in ("costo_pesos", "precio_venta_sugerido"):
                sheet.cell(row, column_by_field[field]).number_format = '$#,##0 "COP"'
    sheet.freeze_panes = "A2"
    for field, values in config["choices"].items():
        column = next(index for index, (name, _label) in enumerate(config["columns"], 1) if name == field)
        validation = DataValidation(type="list", formula1='"' + ",".join(values) + '"', allow_blank=False)
        validation.error = "Selecciona un valor de la lista."
        validation.errorTitle = "Valor no permitido"
        sheet.add_data_validation(validation)
        validation.add(f"{chr(64 + column)}2:{chr(64 + column)}5000")
    info = workbook.create_sheet("Instrucciones")
    info.append(["INSTRUCCIONES DE IMPORTACIÓN"])
    info.append(["1. No cambies los nombres de las columnas."])
    info.append(["2. Elimina la fila de ejemplo o reemplázala con tus datos."])
    info.append(["3. Los campos marcados con * son obligatorios."])
    info.append(["4. Deja ID vacío para crear; escribe un ID existente para actualizar."])
    info.append(["5. Guarda el archivo como .xlsx antes de importarlo."])
    info.column_dimensions["A"].width = 85
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def _text(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _normalize_header(value):
    text = unicodedata.normalize("NFKD", _text(value)).encode("ascii", "ignore").decode().lower()
    return re.sub(r"[^a-z0-9]+", " ", text.replace("*", " ")).strip()


HEADER_ALIASES = {
    "id": ("id", "codigo"),
    "anio_inicio": ("ano inicio", "anio inicio", "desde"),
    "anio_fin": ("ano fin", "anio fin", "hasta"),
    "cantidad_disponible": ("cantidad disponible", "cantidad", "stock", "existencias"),
    "costo_dolares": ("costo usd", "costo dolares", "costo en dolares"),
    "costo_pesos": ("costo cop", "costo pesos", "costo en pesos"),
    "precio_venta_sugerido": ("precio de venta", "precio venta", "precio venta sugerido", "precio"),
    "porcentaje_comision": ("comision", "comision %", "porcentaje comision"),
}


def _convert(field, value, config):
    if value in (None, ""):
        return None
    if field == "activo":
        normalized = _text(value).lower()
        if normalized not in ("sí", "si", "no", "1", "0", "true", "false"):
            raise ValueError("debe ser Sí o No")
        return normalized in ("sí", "si", "1", "true")
    if field in config["integer"] or field == "id":
        return int(value)
    if field in config["decimal"]:
        return Decimal(str(value).replace(",", "."))
    return _text(value)


def import_excel(module, uploaded_file):
    config = IMPORT_CONFIG[module]
    try:
        workbook = load_workbook(uploaded_file, data_only=True, read_only=True)
    except Exception as error:
        return 0, [f"No fue posible abrir el archivo Excel: {error}"]
    sheet = workbook.worksheets[0]
    received = [_normalize_header(cell.value) for cell in sheet[1]]
    received_positions = {header: index for index, header in enumerate(received) if header}
    field_positions = {}
    missing = []
    for field, label in config["columns"]:
        possible_names = {_normalize_header(label), *HEADER_ALIASES.get(field, ())}
        position = next((received_positions[name] for name in possible_names if name in received_positions), None)
        if position is not None:
            field_positions[field] = position
        elif field in config["required"]:
            missing.append(label.replace(" *", ""))
    if missing:
        return 0, ["Faltan estas columnas obligatorias: " + ", ".join(missing) + ". Los espacios, tildes y el orden de las demás columnas no afectan la importación."]

    pending, errors = [], []
    for row_number, row in enumerate(sheet.iter_rows(min_row=2, max_col=sheet.max_column, values_only=True), 2):
        if not any(value not in (None, "") for value in row):
            continue
        raw = {field: row[position] if position < len(row) else None for field, position in field_positions.items()}
        values = {}
        row_errors = []
        for field, label in config["columns"]:
            value = raw.get(field)
            if field in config["required"] and value in (None, ""):
                row_errors.append(f"{label.replace(' *', '')} es obligatorio")
                continue
            try:
                values[field] = _convert(field, value, config)
            except (ValueError, TypeError, InvalidOperation):
                row_errors.append(f"{label.replace(' *', '')} tiene un valor inválido")
        for field, default in config.get("defaults", {}).items():
            if values.get(field) is None:
                values[field] = default
        for field, choices in config["choices"].items():
            if field != "activo" and values.get(field) not in choices:
                row_errors.append(f"{dict(config['columns'])[field].replace(' *', '')} debe ser: {', '.join(choices)}")
        if module == "productos":
            if (values.get("cantidad_disponible") or 0) < 0:
                row_errors.append("Cantidad disponible no puede ser negativa")
            if values.get("anio_inicio") and values.get("anio_fin") and values["anio_fin"] < values["anio_inicio"]:
                row_errors.append("Año fin no puede ser menor que año inicio")
        if module == "vendedores" and values.get("porcentaje_comision") is not None and not 0 <= values["porcentaje_comision"] <= 100:
            row_errors.append("Comisión debe estar entre 0 y 100")
        instance_id = values.pop("id", None)
        if instance_id:
            try:
                instance = config["model"].objects.get(pk=instance_id)
            except config["model"].DoesNotExist:
                # Es común reimportar una hoja después de limpiar la tabla.
                # En ese caso el ID anterior funciona como referencia histórica,
                # pero la fila debe crear un registro nuevo.
                instance = config["model"]()
        else:
            instance = config["model"]()
        if instance:
            for field, value in values.items():
                setattr(instance, field, value)
            try:
                instance.full_clean(exclude=("creado_en", "actualizado_en"))
            except ValidationError as error:
                row_errors.extend(error.messages)
        if row_errors:
            errors.append(f"Fila {row_number}: " + "; ".join(row_errors))
        else:
            pending.append(instance)
    if errors:
        return 0, errors
    with transaction.atomic():
        for instance in pending:
            instance.save()
    return len(pending), []

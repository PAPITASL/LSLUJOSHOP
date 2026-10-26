from decimal import Decimal
from io import BytesIO, StringIO
from pathlib import Path
from tempfile import TemporaryDirectory

from django.apps import apps
from django.contrib.auth.models import User
from django.core.exceptions import ValidationError
from django.core.management import call_command
from django.db import connection, models
from django.test import TestCase, RequestFactory
from django.utils import timezone
from django.urls import reverse
from openpyxl import load_workbook

from .excel_import import build_template, import_excel
from .forms import ClienteForm
from .models import Clientes, Ordenes
from .views import MODULES, _apply_module_filters


class ClientTests(TestCase):
    @classmethod
    def setUpClass(cls):
        if connection.vendor != "sqlite":
            raise RuntimeError("Run with --settings=config.test_settings")
        cls.shop_models = list(apps.get_app_config("gestion").get_models())
        with connection.schema_editor() as editor:
            for model in cls.shop_models:
                editor.create_model(model)
        super().setUpClass()

    @classmethod
    def tearDownClass(cls):
        super().tearDownClass()
        with connection.schema_editor() as editor:
            for model in reversed(cls.shop_models):
                editor.delete_model(model)

    def setUp(self):
        self.customer = Clientes.objects.create(nombre="JOSÉ ROJAS", telefono="314 3119543", ciudad="Bogotá")

    def test_form_and_save_reject_formatted_duplicate_but_allow_edit(self):
        form = ClienteForm(data={"nombre": "  Jose   Rojas ", "telefono": "+57 314-3119543"})
        self.assertFalse(form.is_valid())
        self.assertIn("Ya existe", str(form.errors))
        with self.assertRaises(ValidationError):
            Clientes.objects.create(nombre="Jose Rojas", telefono="3143119543")
        self.customer.ciudad = "Cali"
        self.customer.save()

    def test_shared_name_or_shared_phone_alone_is_allowed(self):
        Clientes.objects.create(nombre="Jose Rojas", telefono="3001234567")
        Clientes.objects.create(nombre="Otra persona", telefono="3143119543")
        Clientes.objects.create(nombre="Sin teléfono", telefono="Sin registrar")
        Clientes.objects.create(nombre="Sin teléfono", telefono="Sin registrar")

    def test_search_ignores_accents_spacing_and_phone_format(self):
        for query in ("jose", "JOSE   ROJAS", "3143119543", "+57 314 3119543", "bogota"):
            with self.subTest(query=query):
                request = RequestFactory().get("/", {"q": query})
                records, _ = _apply_module_filters(Clientes.objects.all(), request, "clientes", MODULES["clientes"])
                self.assertEqual(list(records.values_list("pk", flat=True)), [self.customer.pk])

    def test_excel_duplicate_rows_roll_back_entire_import(self):
        workbook = load_workbook(build_template("clientes"))
        sheet = workbook.worksheets[0]
        sheet.delete_rows(2, sheet.max_row)
        row = [None, "Nuevo cliente", "3001234567", "Cali", "Calle 1", ""]
        sheet.append(row)
        sheet.append(row)
        data = BytesIO()
        workbook.save(data)
        data.seek(0)
        count, errors = import_excel("clientes", data)
        self.assertEqual(count, 0)
        self.assertTrue(errors)
        self.assertEqual(Clientes.objects.count(), 1)

    def test_list_searches_beyond_first_200_and_renders_live_results(self):
        Clientes.objects.bulk_create([Clientes(nombre=f"Otro {i}", telefono=f"300{i:07}") for i in range(201)])
        self.client.force_login(User.objects.create_user(username="clients-test"))
        response = self.client.get(reverse("record_list", args=["clientes"]), {"q": "jose"})
        self.assertEqual(response.status_code, 200)
        self.assertEqual([row.pk for row in response.context["records"]], [self.customer.pk])
        self.assertContains(response, 'id="record-results"')

    def test_merge_preserves_order_and_information_with_backup(self):
        # bulk_create bypasses save to represent pre-existing duplicates.
        duplicate = Clientes(nombre="Jose Rojas", telefono="3143119543", direccion="Calle 9", observaciones="Dato importante")
        Clientes.objects.bulk_create([duplicate])
        values = {field.name: Decimal("0") for field in Ordenes._meta.fields if isinstance(field, models.DecimalField)}
        order = Ordenes.objects.create(numero_orden="TEST", cliente=duplicate, fecha=timezone.now(), estado="Pendiente", **values)
        with TemporaryDirectory() as directory:
            backup = Path(directory) / "backup.json"
            call_command("merge_duplicate_clients", stdout=StringIO())
            self.assertEqual(Clientes.objects.count(), 2)
            call_command("merge_duplicate_clients", apply=True, backup=str(backup), stdout=StringIO())
            self.assertTrue(backup.exists())
        order.refresh_from_db()
        self.customer.refresh_from_db()
        self.assertEqual(order.cliente_id, self.customer.pk)
        self.assertEqual(self.customer.direccion, "Calle 9")
        self.assertIn("Dato importante", self.customer.observaciones)
        self.assertEqual(Clientes.objects.count(), 1)
